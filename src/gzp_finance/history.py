from __future__ import annotations

import csv
import json
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

from sqlalchemy import select

from .db import Classification, HistoricalRecord, Transaction, TransactionComponent
from .domain import BANKS, digest, validate_values
from .rules import normalize_concept


def workbook_payload(source: Path | BytesIO, *, source_name: str | None = None) -> dict:
    """Read the user's original ledger, preserving its row numbers and labels."""
    from openpyxl import load_workbook

    sheet_name = "Transacciones cuentas"
    source_name = source_name or Path(source).name
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except (BadZipFile, OSError, ValueError):
        raise ValueError("No se puede abrir el Excel histórico") from None
    try:
        if sheet_name not in workbook:
            raise ValueError(f"Falta la hoja {sheet_name} en el Excel histórico")
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows()
        headers = [cell.value for cell in next(rows)]
        if len(headers) < 13 or [headers[i] for i in (0, 3, 12)] != ["Fecha", "Banco", "Importe (€)"]:
            raise ValueError("Las columnas del Excel histórico no son las esperadas")
        aliases = {"01. La Caixa": "CaixaBank", "03. MyInvestor": "MyInvestor",
                   "04. Trade Republic": "Trade Republic"}
        targets = ("target_tipo_transaccion", "target_tipo_gasto", "target_fiscalidad",
                   "target_categoria_general", "target_subtipo", "target_activo",
                   "target_detalle", "target_detalle2")
        records = []
        for cells in rows:
            values = [cell.value for cell in cells[:13]]
            if not any(value is not None for value in values):
                continue
            source_row = cells[0].row
            raw_date, bank, raw_amount = values[0], values[3], values[12]
            if not raw_date or not isinstance(bank, str) or not bank.strip():
                raise ValueError(f"Fila {source_row}: fecha o banco ausente")
            if isinstance(raw_date, datetime):
                manual_date = raw_date.date()
            elif isinstance(raw_date, date):
                manual_date = raw_date
            elif isinstance(raw_date, str):
                try:
                    manual_date = datetime.strptime(raw_date.strip(), "%d/%m/%Y").date()
                except ValueError:
                    manual_date = date.fromisoformat(raw_date.strip())
            else:
                raise ValueError(f"Fila {source_row}: fecha inválida")
            if raw_amount is None:
                amount = None
            else:
                numeric = isinstance(raw_amount, (int, float, Decimal)) and not isinstance(raw_amount, bool)
                raw = str(raw_amount).strip()
                if "," in raw:
                    raw = raw.replace(".", "").replace(",", ".")
                try:
                    amount = str(Decimal(raw).quantize(Decimal(".01")))
                except Exception:
                    raise ValueError(f"Fila {source_row}: importe inválido") from None
            record = {"source_row": source_row, "source_bank": aliases.get(bank, bank),
                      "manual_date": manual_date.isoformat(), "amount_eur": amount,
                      "values": validate_values({target: str(values[4 + i]) if values[4 + i] is not None else ""
                                                 for i, target in enumerate(targets)})}
            if raw_amount is not None and not numeric:
                record["excel_amount_is_numeric"] = False
            records.append(record)
        if not records:
            raise ValueError("El Excel histórico no contiene movimientos")
        return {"filename": source_name, "sheet": sheet_name, "records": records}
    finally:
        workbook.close()


def load_workbook(session, source: Path | BytesIO, *, source_name: str | None = None) -> dict:
    """Import or verify the complete original ledger without a bank crosswalk."""
    payload = workbook_payload(source, source_name=source_name)
    identity = {**payload, "records": [
        {k: v for k, v in row.items() if k != "excel_amount_is_numeric"}
        for row in payload["records"]]}
    source_hash = digest(identity)
    existing = {r.source_row: r for r in session.scalars(select(HistoricalRecord))}
    expected_rows = {r["source_row"] for r in payload["records"]}
    if set(existing) - expected_rows or any(r.source_hash != source_hash for r in existing.values()):
        raise ValueError("La base contiene otra versión del histórico; revisa antes de mezclarla")
    inserted = 0
    for row in payload["records"]:
        source_row = row["source_row"]
        amount = Decimal(row["amount_eur"]) if row["amount_eur"] is not None else None
        numeric = row.get("excel_amount_is_numeric", True)
        old = existing.get(source_row)
        if old:
            if (old.source_file != payload["filename"] or old.source_sheet != payload["sheet"]
                    or old.source_bank != row["source_bank"] or old.manual_date != date.fromisoformat(row["manual_date"])
                    or old.amount_eur != amount or old.values != row["values"]
                    or old.excel_amount_is_numeric != numeric):
                raise ValueError(f"Fila {source_row}: la base difiere del Excel original")
            continue
        session.add(HistoricalRecord(source_hash=source_hash, source_file=payload["filename"],
            source_sheet=payload["sheet"], source_row=source_row, source_bank=row["source_bank"],
            manual_date=date.fromisoformat(row["manual_date"]), amount_eur=amount,
            excel_amount_is_numeric=numeric, values=row["values"], bank_input={}))
        inserted += 1
    session.flush()
    summary = validation_summary(session)
    return {"inserted": inserted, "verified": len(payload["records"]),
            "total": summary["excel_total"], "banks": summary["excel_banks"]}


def load_reference(session, path: Path, dataset: Path) -> dict:
    payload = json.loads(path.read_text(encoding="utf-8"))
    # Extraction retains every original row, including banks outside the three importers.
    # Cell typing is supplementary metadata: adding it must not change the identity
    # of an already loaded workbook or duplicate its historical rows.
    identity = {**payload, "records": [
        {k: v for k, v in row.items() if k != "excel_amount_is_numeric"}
        for row in payload["records"]]}
    source_hash = digest(identity)
    with dataset.open(encoding="utf-8-sig", newline="") as f:
        bank_rows = {(r["source_bank"], int(r["source_row"])): r for r in csv.DictReader(f)}
    existing = {r.source_row: r for r in session.scalars(select(HistoricalRecord).where(
        HistoricalRecord.source_hash == source_hash))}
    other_source = session.scalar(select(HistoricalRecord.id).where(HistoricalRecord.source_hash != source_hash))
    if other_source:
        raise ValueError("Ya existe otro Excel de referencia; no mezcles versiones del mismo histórico")
    inserted = 0
    for row in payload["records"]:
        numeric = row.get("excel_amount_is_numeric", True)
        if not isinstance(numeric, bool):
            raise ValueError("excel_amount_is_numeric debe ser booleano")
        if row["source_row"] in existing:
            if "excel_amount_is_numeric" in row:
                existing[row["source_row"]].excel_amount_is_numeric = numeric
            continue
        original = validate_values(row["values"])
        matched = bank_rows.get((row["source_bank"], row["source_row"]), {})
        inputs = {k: v for k, v in matched.items() if not k.startswith("target_")
                  and k not in {"manual_date", "model_text"}}
        session.add(HistoricalRecord(source_hash=source_hash, source_file=payload["filename"],
            source_sheet=payload["sheet"], source_row=row["source_row"], source_bank=row["source_bank"],
            manual_date=date.fromisoformat(row["manual_date"]),
            amount_eur=Decimal(row["amount_eur"]) if row["amount_eur"] is not None else None,
            excel_amount_is_numeric=numeric,
            values=original, bank_input=inputs, match_confidence=matched.get("match_confidence", "")))
        inserted += 1
    session.flush()
    return {"inserted": inserted, **link_history(session)}


def link_history(session) -> dict:
    transactions = list(session.scalars(select(Transaction)))
    by_id, by_key = defaultdict(list), defaultdict(list)
    for tx in transactions:
        by_key[(tx.source_bank, str(tx.bank_date), tx.amount_eur, tx.concept_key)].append(tx)
        if tx.transaction_id:
            by_id[(tx.source_bank, tx.transaction_id)].append(tx)
    owners = {t.id: t for t in transactions}
    for component in session.scalars(select(TransactionComponent)):
        key = (component.source_bank, component.external_id)
        owner = owners[component.transaction_id]
        if owner not in by_id[key]:
            by_id[key].append(owner)
    matched, unresolved = 0, 0
    for ref in session.scalars(select(HistoricalRecord).order_by(HistoricalRecord.source_row)):
        if ref.transaction_id:
            matched += 1
            continue
        row = ref.bank_input
        if not row:
            unresolved += 1
            continue
        candidates = by_id.get((ref.source_bank, row.get("transaction_id")), []) if row.get("transaction_id") else []
        amount = Decimal(row["amount_eur"]).quantize(Decimal(".01"))
        if not candidates:
            candidates = by_key.get((ref.source_bank, row["bank_date"], amount,
                                     normalize_concept(row.get("concept_raw", ""))), [])
            if row.get("account_ref"):
                candidates = [t for t in candidates if t.account_ref == row["account_ref"]]
        if len(candidates) == 1:
            ref.transaction_id = candidates[0].id
            ref.linkage_status = "matched"
            matched += 1
        else:
            ref.linkage_status = "ambiguous" if candidates else "unmatched"
            unresolved += 1
    session.flush()
    return {"matched": matched, "unresolved": unresolved}


def totals(amounts) -> dict:
    amounts = list(amounts)
    known = [Decimal(a) for a in amounts if a is not None]
    incoming = sum((a for a in known if a > 0), Decimal(0))
    outgoing = -sum((a for a in known if a < 0), Decimal(0))
    return {"count": len(amounts), "missing_amount": len(amounts) - len(known),
            "income": str(incoming.quantize(Decimal('.01'))),
            "expenses": str(outgoing.quantize(Decimal('.01'))),
            "net": str((incoming - outgoing).quantize(Decimal('.01')))}


def validation_summary(session) -> dict:
    txs = list(session.scalars(select(Transaction)))
    refs = list(session.scalars(select(HistoricalRecord)))
    confirmed = set(session.scalars(select(Classification.transaction_id).where(
        Classification.confirmed_by_user.is_(True))))
    by_id = {t.id: t for t in txs}
    linked_ids = {r.transaction_id for r in refs if r.transaction_id}
    banks = []
    excel_banks = []
    source_labels = {"CaixaBank": "01. La Caixa", "MyInvestor": "03. MyInvestor",
                     "Trade Republic": "04. Trade Republic"}
    for bank in sorted({r.source_bank for r in refs}, key=lambda name: source_labels.get(name, name)):
        original = [r for r in refs if r.source_bank == bank]
        excel_banks.append({"bank": source_labels.get(bank, bank), "totals": totals(
            r.amount_eur if r.excel_amount_is_numeric else None for r in original)})
    for bank in BANKS:
        imported = [t for t in txs if t.source_bank == bank]
        original = [r for r in refs if r.source_bank == bank]
        imported_dates = [t.bank_date for t in imported]
        original_dates = [r.manual_date for r in original]
        # One original row per bank movement; grouped/ambiguous relations are reported separately.
        links = defaultdict(list)
        for r in original:
            if r.transaction_id:
                links[r.transaction_id].append(r)
        comparable = [rr[0] for tid, rr in links.items() if len(rr) == 1 and
                      rr[0].amount_eur is not None and rr[0].excel_amount_is_numeric and tid in by_id]
        comparable_ids = {r.id for r in comparable}
        excel = totals(r.amount_eur for r in comparable)
        actual = totals(by_id[r.transaction_id].amount_eur for r in comparable)
        deltas = {k: str((Decimal(actual[k]) - Decimal(excel[k])).quantize(Decimal('.01')))
                  for k in ("income", "expenses", "net")}
        mismatched = sum(by_id[r.transaction_id].amount_eur != r.amount_eur for r in comparable)
        banks.append({"bank": bank, "imported": totals(t.amount_eur for t in imported),
            "imported_dates": {"first": str(min(imported_dates)) if imported_dates else None,
                               "last": str(max(imported_dates)) if imported_dates else None},
            "excel_dates": {"first": str(min(original_dates)) if original_dates else None,
                            "last": str(max(original_dates)) if original_dates else None},
            "new": totals(t.amount_eur for t in imported if t.id not in linked_ids),
            "historical": totals(t.amount_eur for t in imported if t.id in linked_ids),
            "excel": totals(r.amount_eur if r.excel_amount_is_numeric else None for r in original),
            "outside_comparison": totals(r.amount_eur if r.excel_amount_is_numeric else None
                                         for r in original if r.id not in comparable_ids),
            "text_amount_rows": sum(r.amount_eur is not None and not r.excel_amount_is_numeric
                                    for r in original),
            "linked_rows": sum(len(v) for v in links.values()),
            "unlinked_rows": sum(not r.transaction_id for r in original),
            "ambiguous_rows": sum(r.linkage_status == "ambiguous" for r in original),
            "grouped_rows": sum(len(rr) for rr in links.values() if len(rr) > 1),
            "confirmed": sum(t.id in confirmed for t in imported),
            "comparison": {"excel": excel, "bank": actual, "delta": deltas,
                           "mismatched_rows": mismatched, "matches": bool(comparable) and mismatched == 0}})
    return {"banks": banks, "excel_banks": excel_banks,
            "excel_total": totals(r.amount_eur if r.excel_amount_is_numeric else None for r in refs),
            "excel_records": len(refs), "imported_records": len(txs),
            "historical_records": len(linked_ids), "new_records": len(txs) - len(linked_ids),
            "confirmed_records": len(confirmed)}
