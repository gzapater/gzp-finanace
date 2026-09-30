import csv
import json
import re
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select

from gzp_finance.db import Classification, HistoricalRecord, TrainingExample, Transaction
from gzp_finance.history import link_history, load_reference, load_workbook, validation_summary
from gzp_finance.service import import_statement
from test_importers import MI, TR
from test_web import client


def reference(tmp_path, *, amount='-1299.87', confidence='Alta'):
    p = tmp_path / 'reference.json'
    p.write_text(json.dumps({'filename':'example.xlsx','sheet':'Transactions','records':[
        {'source_row':2,'source_bank':'MyInvestor','manual_date':'2026-03-31',
         'amount_eur':amount,'values':{'target_subtipo':'Original label'}}]}))
    d = tmp_path / 'dataset.csv'
    with d.open('w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=['source_bank','source_row','bank_date','concept_raw','amount_eur','match_confidence'])
        w.writeheader(); w.writerow({'source_bank':'MyInvestor','source_row':2,'bank_date':'2026-04-01',
            'concept_raw':'Fondo ejemplo','amount_eur':'-1299.87','match_confidence':confidence})
    return p,d


def original_workbook(tmp_path):
    path=tmp_path/'original.xlsx'
    workbook=Workbook();sheet=workbook.active;sheet.title='Transacciones cuentas'
    sheet.append(['Fecha','mes','Año','Banco','Tipo transacción','Tipo de gasto','Fiscalidad',
                  'Categoria General','Subtipo','Activo','Detalle','Detalle 2','Importe (€)'])
    sheet.append([date(2026,4,1),4,2026,'01. La Caixa','Ingreso','','','Ingresos','Nómina',
                  'Trabajo','','',32.40])
    sheet.append(['02/04/2026',4,2026,'03. MyInvestor','Inversión','','','Inversión','Fondo',
                  'Activo','','','0,02'])
    workbook.save(path)
    return path


def test_original_workbook_is_canonical_ledger_and_reimport_is_checked(session,tmp_path):
    path=original_workbook(tmp_path)
    first=load_workbook(session,path)
    assert first['inserted']==2 and first['verified']==2
    assert first['total']['net']=='32.40'
    assert [(b['bank'],b['totals']['net']) for b in first['banks']] == [
        ('01. La Caixa','32.40'),('03. MyInvestor','0.00')]
    assert load_workbook(session,path)['inserted']==0
    assert session.scalar(select(func.count()).select_from(HistoricalRecord))==2
    assert session.scalar(select(func.count()).select_from(Transaction))==0
    assert session.scalar(select(func.count()).select_from(TrainingExample))==0
    from openpyxl import load_workbook as open_workbook
    changed=open_workbook(path);changed.active['M2']=33.40;changed.save(path)
    with pytest.raises(ValueError,match='otra versión'):
        load_workbook(session,path)
    assert validation_summary(session)['excel_total']['net']=='32.40'


def test_original_workbook_upload_shows_filtered_ledger_total(engine,tmp_path):
    path=original_workbook(tmp_path)
    c=client(engine);c.auth=('test','synthetic-test-password')
    html=c.get('/').text
    token=re.search(r'name="csrf-token" content="([^"]+)"',html)[1]
    c.headers['x-csrf-token']=token
    with path.open('rb') as file:
        uploaded=c.post('/api/history/import',files={'file':('original.xlsx',file)})
    assert uploaded.status_code==200 and uploaded.json()['verified']==2
    assert uploaded.json()['total']['net']=='32.40'
    assert c.get('/api/history?bank=CaixaBank').json()['totals']['net']=='32.40'
    assert c.get('/api/history?bank=MyInvestor').json()['totals']['net']=='0.00'
    assert c.get('/api/validation').json()['excel_total']['net']=='32.40'
    assert '01. La Caixa' in c.get('/').text


def test_reference_before_import_links_later_and_never_confirms_or_trains(session,tmp_path):
    p,d=reference(tmp_path,confidence='Media')
    assert load_reference(session,p,d)['inserted'] == 1
    assert load_reference(session,p,d)['inserted'] == 0
    import_statement(session,'MyInvestor','a.csv',MI.encode())
    h=session.scalar(select(HistoricalRecord))
    assert h.transaction_id and h.match_confidence == 'Media'
    assert not session.scalar(select(Classification)).confirmed_by_user
    assert session.scalar(select(Classification)).values == {}
    assert session.scalar(select(func.count()).select_from(TrainingExample)) == 0
    stats=validation_summary(session)
    assert stats['imported_records'] == 1 and stats['new_records'] == 0
    assert stats['banks'][2]['comparison']['matches']
    assert stats['banks'][2]['excel_dates']['first'] == '2026-03-31'
    assert stats['banks'][2]['imported_dates']['first'] == '2026-04-01'


def test_duplicate_payments_are_ambiguous_not_arbitrarily_linked(session,tmp_path):
    import_statement(session,'MyInvestor','a.csv',(MI+MI.splitlines()[1]+'\n').encode())
    p,d=reference(tmp_path)
    assert load_reference(session,p,d)['matched'] == 0
    assert session.scalar(select(HistoricalRecord)).linkage_status == 'ambiguous'
    stats=validation_summary(session)['banks'][2]
    assert stats['new']['count'] == 2
    assert stats['comparison']['matches'] is False
    assert stats['ambiguous_rows'] == 1


def test_grouped_fills_not_double_counted_or_compared_as_single_excel_rows(session):
    import_statement(session,'Trade Republic','a.csv',TR.encode())
    for n,external_id,amount in [(2,'id-a','-100.00'),(3,'id-b','-50.00')]:
        session.add(HistoricalRecord(source_hash='a',source_file='a.xlsx',source_sheet='x',source_row=n,
            source_bank='Trade Republic',manual_date=date(2026,4,1),amount_eur=Decimal(amount),values={},
            bank_input={'transaction_id':external_id,'bank_date':'2026-04-01','amount_eur':amount}))
    session.flush(); assert link_history(session)['matched'] == 2
    stats=validation_summary(session)['banks'][0]
    assert stats['historical']['count'] == 1
    assert stats['grouped_rows'] == 2
    assert stats['comparison']['excel']['count'] == 0
    assert stats['imported']['net'] == '-150.00'


def test_amount_difference_and_missing_source_are_visible(session,tmp_path):
    import_statement(session,'MyInvestor','a.csv',MI.encode())
    p,d=reference(tmp_path,amount='-1300.00')
    load_reference(session,p,d)
    c=validation_summary(session)['banks'][2]['comparison']
    assert c['delta']['net'] == '0.13' and not c['matches']
    h=session.scalar(select(HistoricalRecord)); h.amount_eur=None; session.flush()
    summary=validation_summary(session)['banks'][2]
    assert summary['excel']['missing_amount'] == 1
    assert not summary['comparison']['matches']


def test_excel_text_amount_does_not_enter_pivot_kpis_but_keeps_bank_match(session,tmp_path):
    import_statement(session,'MyInvestor','a.csv',MI.encode())
    p,d=reference(tmp_path)
    assert load_reference(session,p,d)['inserted'] == 1
    h=session.scalar(select(HistoricalRecord))
    original_hash=h.source_hash
    payload=json.loads(p.read_text())
    payload['records'][0]['excel_amount_is_numeric']=False
    p.write_text(json.dumps(payload))
    assert load_reference(session,p,d)['inserted'] == 0
    assert h.source_hash == original_hash and h.transaction_id
    assert not h.excel_amount_is_numeric
    stats=validation_summary(session)
    assert stats['excel_total']['net'] == '0.00'
    assert stats['excel_total']['missing_amount'] == 1
    bank=stats['banks'][2]
    assert bank['text_amount_rows'] == 1
    assert bank['comparison']['excel']['count'] == 0
    assert bank['outside_comparison']['count'] == 1


def test_history_api_and_origin_filters_are_authenticated_and_keep_reference_labels(engine,tmp_path):
    from sqlalchemy.orm import Session
    with Session(engine) as s:
        import_statement(s,'MyInvestor','a.csv',MI.encode());p,d=reference(tmp_path);load_reference(s,p,d);s.commit()
    c=client(engine)
    assert c.get('/api/history').status_code == 401
    assert c.get('/api/validation').status_code == 401
    c.auth=('test','synthetic-test-password')
    assert c.get('/api/transactions?origin=new').json()['total'] == 0
    assert c.get('/api/transactions?origin=historical').json()['total'] == 1
    assert c.get('/api/history').json()['items'][0]['values']['target_subtipo'] == 'Original label'
    assert 'Original label' in c.get('/api/vocabulary').json()['fields']['target_subtipo']
    assert c.get('/api/transactions?origin=invalid').status_code == 422
    assert c.get('/api/history?offset=-1').status_code == 422
    assert c.get('/api/validation').json()['excel_records'] == 1
